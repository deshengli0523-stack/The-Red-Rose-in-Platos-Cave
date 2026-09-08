from __future__ import annotations

import importlib
from pathlib import Path

import pytest

from consultation_kb.knowledge.extractors import DocumentExtractor, ExtractionError


FIXTURES = Path(__file__).resolve().parents[2] / "fixtures" / "consultation_kb" / "sources"


@pytest.mark.parametrize(
    ("name", "document_type", "locator_kind"),
    [
        ("synthetic_classic.md", "md", "source_line_span"),
        ("synthetic_counseling.md", "md", "source_line_span"),
        ("synthetic_career.csv", "csv", "source_table_span"),
    ],
)
def test_local_text_extractors_have_canonical_coordinates(
    name: str, document_type: str, locator_kind: str
) -> None:
    result = DocumentExtractor().extract(FIXTURES / name)
    assert result.status == "ok"
    assert result.document_type == document_type
    assert result.blocks
    assert all(block.extractor_version for block in result.blocks)
    assert all(block.locator.locator_kind == locator_kind for block in result.blocks)
    assert all("pages:" not in block.locator.display_locator for block in result.blocks)


def test_docx_and_xlsx_are_generated_locally_with_exact_coordinates(tmp_path: Path) -> None:
    from docx import Document
    from openpyxl import Workbook  # type: ignore[import-untyped]

    docx_path = tmp_path / "synthetic.docx"
    document = Document()
    document.add_paragraph("Synthetic paragraph")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    document.save(str(docx_path))

    xlsx_path = tmp_path / "synthetic.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["B2"] = "first"
    sheet["C2"] = "second"
    workbook.save(xlsx_path)

    docx_result = DocumentExtractor().extract(docx_path)
    xlsx_result = DocumentExtractor().extract(xlsx_path)
    assert {block.locator.locator_kind for block in docx_result.blocks} == {
        "source_paragraph_span",
        "source_table_span",
    }
    assert xlsx_result.blocks[0].locator.display_locator == "sheet:1;range:B2-C2"


def test_programmatically_generated_pdf_extracts_text_and_page_block(tmp_path: Path) -> None:
    from pypdf import PdfWriter
    from pypdf.generic import (
        DecodedStreamObject,
        DictionaryObject,
        NameObject,
    )

    path = tmp_path / "synthetic.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 40 250 Td (Synthetic evidence) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    with path.open("wb") as stream:
        writer.write(stream)

    result = DocumentExtractor().extract(path)
    assert result.status == "ok"
    assert "Synthetic evidence" in result.blocks[0].text
    assert result.blocks[0].locator.display_locator == "pages:1:1-1:1"


def test_scanned_pdf_is_quarantined_instead_of_fake_success(tmp_path: Path) -> None:
    from pypdf import PdfWriter

    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with path.open("wb") as stream:
        writer.write(stream)

    result = DocumentExtractor().extract(path)
    assert result.status == "quarantined"
    assert result.blocks == ()
    assert result.quarantine_reason == "SCANNED_PDF_OCR_REQUIRED"


def test_missing_optional_extractor_is_explicit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "synthetic.pdf"
    path.write_bytes(b"%PDF-synthetic")
    original = importlib.import_module

    def missing(name: str, package: str | None = None):  # type: ignore[no-untyped-def]
        if name == "pypdf":
            raise ImportError(name)
        return original(name, package)

    monkeypatch.setattr(importlib, "import_module", missing)
    with pytest.raises(ExtractionError, match="OPTIONAL_EXTRACTOR_MISSING"):
        DocumentExtractor().extract(path)
